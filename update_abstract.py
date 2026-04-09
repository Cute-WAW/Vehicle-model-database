# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Title
title = doc.add_heading('基于规则引擎与LightGBM的二手车残值估价系统设计', level=1)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Abstract Heading (Chinese)
abstract_heading_zh = doc.add_heading('摘要', level=2)
abstract_heading_zh.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Paragraph 1
p1 = doc.add_paragraph()
p1.add_run('随着我国二手车市场的快速发展，人们对二手车的精准定价和售卖需求日益突出。然而，由于二手车具有“一车一况、一车一价”的非标准化特性，导致市场信息不对称、定价缺乏透明度与科学性。')
p1.paragraph_format.first_line_indent = Pt(21)

# Paragraph 2
p2 = doc.add_paragraph()
p2.add_run('针对上述问题，本文设计并实现了一套基于规则引擎与LightGBM的二手车残值估价系统。该系统所使用的真实交易数据来源于车易拍、懂车帝等二手车交易平台。在数据预处理与车型对齐阶段，本文设计了一种基于专家规则引擎的车型实体提取与匹配算法。通过构建包含品牌、车系、年款、排量等多维度配置的知识库，结合自定义实体抽取规则与多级相似度计算模型，实现了海量非标准车辆描述到标准车型库的精准映射。在价格预测模块，本文引入了LightGBM机器学习算法，深度挖掘车龄、行驶里程、品牌保值率、城市级别及车辆配置等多维特征。充分利用LightGBM在处理大规模高维数据时的高效性以及捕捉复杂非线性衰减规律的优势，完成了残值率预测模型的训练。系统上线后，能够根据用户输入的车辆多维条件综合评估残值率，从而实现对二手车价格的精准预测。最后，本文采用模块化架构开发了完整的估价应用以及Web测试界面。')
p2.paragraph_format.first_line_indent = Pt(21)

# Keywords
p_kw = doc.add_paragraph()
kw_bold = p_kw.add_run('关键词：')
kw_bold.bold = True
p_kw.add_run('二手车估价；残值模型；规则引擎；车型匹配；LightGBM；特征工程')

# Abstract Heading (English)
abstract_heading_en = doc.add_heading('Abstract', level=2)
abstract_heading_en.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Paragraph 1 (English)
p1_en = doc.add_paragraph()
p1_en.add_run('With the rapid development of the used car market in China, the demand for accurate pricing and selling of used cars has become increasingly prominent. However, due to the non-standardized characteristics of used cars—often described as "one condition per car, one price per car"—the market suffers from information asymmetry and a lack of transparency and scientific rigor in pricing.')
p1_en.paragraph_format.first_line_indent = Pt(21)

# Paragraph 2 (English)
p2_en = doc.add_paragraph()
p2_en.add_run('To address these issues, this paper designs and implements a used car residual value estimation system based on a rule engine and LightGBM. The real transaction data used in this system is sourced from used car trading platforms such as Cheyipai and Dongchedi. In the data preprocessing and vehicle alignment phase, an expert rule engine-based entity extraction and matching algorithm is designed. By constructing a knowledge base containing multi-dimensional configurations such as brand, series, model year, and engine displacement, combined with custom entity extraction rules and a multi-level similarity calculation model, the precise mapping of massive non-standard vehicle descriptions to a standard vehicle database is achieved. In the price prediction module, the LightGBM machine learning algorithm is introduced to deeply mine multi-dimensional features, including vehicle age, mileage, brand retention rate, city tier, and vehicle configurations. By fully leveraging LightGBM\'s computational efficiency in handling large-scale, high-dimensional data and its advantage in capturing complex non-linear depreciation patterns, the residual value prediction model is successfully trained. Once deployed, the system can comprehensively evaluate the residual value rate based on multiple vehicle conditions input by users, thereby achieving accurate predictions of used car prices. Finally, a complete estimation application and a Web test interface are developed using a modular architecture.')
p2_en.paragraph_format.first_line_indent = Pt(21)

# Keywords (English)
p_kw_en = doc.add_paragraph()
kw_bold_en = p_kw_en.add_run('Key Words: ')
kw_bold_en.bold = True
p_kw_en.add_run('Used Car Pricing; Residual Value Model; Rule Engine; Vehicle Matching; LightGBM; Feature Engineering')

doc.save(r'D:\BaiduNetdiskDownload\车型库映射\车型库映射\doc\毕设\基于规则引擎与LightGBM的二手车残值估价系统设计_摘要.docx')
print("Successfully generated updated Word document.")
