
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_patent_doc():
    doc = Document()
    
    # Title
    title = doc.add_heading('一种基于多源特征检索与动态置信度调整的二手车残值预测方法', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Abstract
    doc.add_heading('摘要', level=1)
    abstract_text = (
        "本发明公开了一种基于多源特征检索与动态置信度调整的二手车残值预测方法。该方法包括：构建包含多维车辆特征的成交数据索引；"
        "基于分级特征匹配策略从成交数据中推断新车价格；利用品牌车系模型或车辆类别模型进行基础残值预测；"
        "检索相近车辆并计算基于特征相似度的置信度得分；根据置信度得分动态调整基础预测结果；最后结合C2B2C模型生成多场景价格矩阵。"
        "本发明有效解决了稀疏样本预测偏差大、新车价格缺失及配置差异导致的估值不准问题，提高了二手车估值的准确性和鲁棒性。"
    )
    doc.add_paragraph(abstract_text)
    
    # Field of Invention
    doc.add_heading('技术领域', level=1)
    doc.add_paragraph(
        "本发明涉及汽车数据处理与价值评估技术领域，具体涉及一种基于多源特征检索与动态置信度调整的二手车残值预测方法。"
    )
    
    # Background
    doc.add_heading('背景技术', level=1)
    doc.add_paragraph(
        "二手车“一车一况、一车一价”的非标品特性使得价格预测具有极大的挑战性。现有技术通常采用单一的统计模型或机器学习回归模型进行预测。"
        "然而，单纯的统计模型难以处理稀疏数据，而回归模型在面对特定配置（如不同排量、变速箱、版式）的差异时，往往因特征粒度不足而导致预测偏差。"
        "此外，新车价格作为残值计算的重要基准，常常因历史数据缺失或款型匹配错误而导致计算基础不准确。现有技术缺乏一种能够结合历史成交数据检索与模型预测优势，"
        "并能根据相似样本的可信度动态调整预测结果的混合方法。"
    )
    
    # Summary of Invention
    doc.add_heading('发明内容', level=1)
    doc.add_paragraph(
        "本发明的目的在于提供一种基于多源特征检索与动态置信度调整的二手车残值预测方法，以解决现有技术中存在的预测精度不足、对特定配置差异不敏感等问题。"
    )
    doc.add_paragraph("本发明提供如下技术方案：")
    
    claims = [
        "一种基于多源特征检索与动态置信度调整的二手车残值预测方法，其特征在于，包括以下步骤：",
        "步骤S1：构建成交数据索引。基于历史成交数据，提取车辆全称中的实体特征（包括年款、排量、变速箱、版式），构建支持多维度检索的倒排索引库；",
        "步骤S2：新车价格推断。接收待预测车辆信息，在索引库中检索相近车辆，采用分级特征匹配策略（优先全特征匹配，逐级降级）筛选有效样本，计算其新车价格中位数作为待预测车辆的基准新车价格；",
        "步骤S3：基础模型预测。优先调用针对该品牌车系训练的回归模型进行价格预测；若不存在，则回退调用针对该车辆类别（如紧凑型车-合资）的通用回归模型，得到基础预测价格；",
        "步骤S4：相近车辆检索与置信度计算。基于待预测车辆的特征（年限、里程、评级、城市及实体特征），从索引库中检索Top-K个相近车辆，并根据特征匹配程度计算综合相似度得分；",
        "步骤S5：动态置信度调整。计算相近车辆的加权平均价格，结合相似度得分计算置信度系数，利用该系数对基础预测价格进行动态加权调整，得到最终预测价格；",
        "步骤S6：多场景价格生成。基于C2B2C模型将最终预测价格映射为包含B2B、C2B、C2C等不同交易场景的价格矩阵。"
    ]
    
    for claim in claims:
        p = doc.add_paragraph(claim)
        p.style = 'List Paragraph'
    
    # Detailed Description
    doc.add_heading('具体实施方式', level=1)
    
    sections = [
        ("1. 数据索引构建模块", 
         "系统首先加载历史成交数据（merged_residual_value_data.csv），利用实体抽取器（EntityExtractor）从车辆全称中解析出年款（Year）、排量（Displacement）、变速箱（Transmission）和版式（Trim）等关键特征。构建品牌车系索引、品牌索引及车辆类别索引，并缓存为pkl文件（residual_data_index.pkl）以加速检索。"),
         
        ("2. 新车价格智能推断策略", 
         "为解决新车价格缺失或不准的问题，本发明提出五级匹配策略（_infer_new_price）：\n"
         "- Level 1：年款+排量+变速箱+版式完全匹配；\n"
         "- Level 2：年款+排量+变速箱匹配；\n"
         "- Level 3：年款+排量匹配；\n"
         "- Level 4：仅年款匹配；\n"
         "- Level 5：使用Top-5相近车辆兜底。\n"
         "系统按优先级顺序查找样本，一旦找到即计算中位数返回，确保基准价格的精准性。"),
         
        ("3. 双层模型预测架构", 
         "采用“品牌车系-车辆类别”双层兜底架构（PricePredictor）：\n"
         "优先使用针对特定品牌车系（如“大众-POLO”）训练的高精度模型；"
         "若该车系样本不足，自动降级使用车辆类别（如“轿车-小型车-合资”）的通用模型，保证预测覆盖率100%。"),
         
        ("4. 基于置信度的动态调整算法", 
         "系统计算检索到的相近车辆的平均校正价格（Similar_Avg）。定义置信度（Confidence）为相近车辆平均相似度分数与满分的比值。\n"
         "调整公式为：P_final = P_model * (1 + Deviation * Confidence * Factor)\n"
         "其中Deviation为(Similar_Avg - P_model)/P_model。该机制使得在相似车辆高度匹配时，预测结果更倾向于市场成交价；在相似车辆匹配度低时，更信赖模型预测值。"),
         
        ("5. 启发式规则修正", 
         "在模型输出基础上，引入专家规则修正：\n"
         "- 新能源车（NEV）：价格上浮3%；\n"
         "- 准新车（3年以内）：价格上浮5%；\n"
         "- 高价车（>15万）：价格上浮5%。")
    ]
    
    for title, content in sections:
        doc.add_heading(title, level=2)
        doc.add_paragraph(content)

    # Save
    output_path = os.path.join(os.getcwd(), 'output', 'patent_draft.docx')
    if not os.path.exists(os.path.dirname(output_path)):
        os.makedirs(os.path.dirname(output_path))
        
    doc.save(output_path)
    print(f"Patent document generated at: {output_path}")

if __name__ == "__main__":
    create_patent_doc()
