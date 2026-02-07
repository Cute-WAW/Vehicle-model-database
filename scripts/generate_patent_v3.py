import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_font(run, font_name='宋体', size=None, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if size:
        run.font.size = size
    if bold:
        run.bold = True

def create_patent_doc_v3():
    doc = Document()
    
    # --- Title ---
    # 标题：一种基于多源特征检索与自适应分段非线性回归的二手车残值预测方法
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run('一种基于多源特征检索与自适应分段非线性回归的二手车残值预测方法')
    set_font(run, '黑体', Pt(16), True)
    
    doc.add_paragraph()  # Spacer

    # --- Abstract ---
    # 摘要
    head_para = doc.add_paragraph()
    run = head_para.add_run('摘要')
    set_font(run, '黑体', Pt(14), True)
    
    abstract_text = (
        "本发明公开了一种基于多源特征检索与自适应分段非线性回归的二手车残值预测方法。该方法首先通过实体提取技术构建包含年款、排量、变速箱及版式的多维车辆特征索引；"
        "针对新车价格缺失问题，提出五级特征匹配推断策略；"
        "在预测阶段，采用“分段-整体-类别”三级竞争优选机制，针对不同车龄段（青年/中年/老年）自动匹配最佳非线性回归模型（指数、多项式或幂函数）；"
        "进而利用基于特征相似度的动态置信度算法，融合相近车辆的市场成交价对模型预测结果进行修正；"
        "最后结合C2B2C模型生成多场景价格矩阵。本发明有效解决了传统LightGBM等机器学习模型在小样本长尾车型上过拟合、非线性衰减拟合不足及新车基准缺失的问题，显著提升了二手车估值的准确性与鲁棒性。"
    )
    p = doc.add_paragraph(abstract_text)
    p.paragraph_format.first_line_indent = Inches(0.3)
    set_font(p.runs[0], '宋体', Pt(12))

    doc.add_paragraph()  # Spacer

    # --- Technical Field ---
    # 一、技术领域
    head_para = doc.add_paragraph()
    run = head_para.add_run('一、技术领域')
    set_font(run, '黑体', Pt(14), True)
    
    p = doc.add_paragraph(
        "本发明涉及汽车数据挖掘与价值评估技术领域，具体涉及一种基于多源特征检索与自适应分段非线性回归的二手车残值预测方法。"
    )
    p.paragraph_format.first_line_indent = Inches(0.3)
    set_font(p.runs[0], '宋体', Pt(12))

    # --- Background ---
    # 二、背景技术
    head_para = doc.add_paragraph()
    run = head_para.add_run('二、背景技术')
    set_font(run, '黑体', Pt(14), True)
    
    background_text = (
        "二手车作为一种典型的非标品，其价格受品牌、车系、车龄、里程、车况以及配置（如排量、变速箱、版式）等多种因素影响，具有“一车一况、一车一价”的特征。"
        "现有估值方法主要存在以下痛点：\n"
        "1. 新车价格基准缺失：二手车残值率的计算依赖于准确的新车指导价，但历史车型库数据常存在缺失或匹配错误，导致估值基准偏移。\n"
        "2. 传统模型局限性：简单的线性回归难以拟合二手车价格随车龄呈现的“早期陡峭、后期平缓”的非线性衰减规律；而复杂的集成树模型（如LightGBM、XGBoost）虽然拟合能力强，但在样本稀疏的长尾车型（如冷门车系）上极易过拟合，且缺乏可解释性。\n"
        "3. 配置差异被忽略：同一车系不同配置（如手动vs自动）价格差异巨大，通用模型往往难以捕捉这些细微特征带来的溢价。\n"
        "4. 缺乏动态修正机制：模型预测通常是静态的，难以结合实时检索到的高相似度市场成交案例进行动态校准。"
    )
    p = doc.add_paragraph(background_text)
    p.paragraph_format.first_line_indent = Inches(0.3)
    set_font(p.runs[0], '宋体', Pt(12))

    # --- Summary of Invention ---
    # 三、发明内容
    head_para = doc.add_paragraph()
    run = head_para.add_run('三、发明内容')
    set_font(run, '黑体', Pt(14), True)
    
    p = doc.add_paragraph(
        "本发明的目的在于提供一种基于多源特征检索与自适应分段非线性回归的二手车残值预测方法，通过多层级的特征匹配与分段模型优选，解决上述技术问题。"
    )
    p.paragraph_format.first_line_indent = Inches(0.3)
    set_font(p.runs[0], '宋体', Pt(12))
    
    p = doc.add_paragraph("本发明提供如下技术方案：")
    set_font(p.runs[0], '宋体', Pt(12))

    claims = [
        "1. 一种基于多源特征检索与自适应分段非线性回归的二手车残值预测方法，其特征在于，包括以下步骤：\n"
        "S1：构建多维成交数据索引。基于历史成交数据，利用实体提取技术解析车辆全称中的年款、排量、变速箱及版式特征，构建倒排索引库；\n"
        "S2：新车价格智能推断。接收待预测车辆信息，采用五级特征匹配策略（优先全特征匹配，逐级降级）从索引库中检索有效样本，计算中位数作为基准新车价格；\n"
        "S3：自适应分段非线性模型预测。根据待预测车辆的车龄，将其划分至青年（0.5-5年）、中年（5-10年）或老年（10-20年）分段，自动调用对应分段的非线性回归模型（指数/多项式/幂函数）进行预测；若分段模型缺失，则触发“分段->整体->类别”三级回退机制；\n"
        "S4：动态置信度修正。检索Top-K相近车辆，计算基于特征匹配度的综合相似度得分及置信度系数，利用该系数对基础预测结果进行加权修正；\n"
        "S5：多场景价格生成。基于C2B2C转换模型，将最终预测值映射为B2B、C2B、C2C等多场景价格矩阵。",
        
        "2. 根据权利要求1所述的方法，其特征在于，步骤S2中的五级特征匹配策略包括：\n"
        "Level 1：年款+排量+变速箱+版式完全匹配；\n"
        "Level 2：年款+排量+变速箱匹配（忽略版式）；\n"
        "Level 3：年款+排量匹配（忽略变速箱）；\n"
        "Level 4：仅年款匹配；\n"
        "Level 5：使用相似度最高的Top-5车辆。\n"
        "系统依次尝试上述级别，一旦获取到足够样本即计算中位数并返回，确保新车价格的精准性与覆盖率。",

        "3. 根据权利要求1所述的方法，其特征在于，步骤S3中的三级回退机制具体为：\n"
        "第一级（分段模型）：优先加载当前品牌车系在特定车龄段（如“本田-飞度_young”）的专用模型；\n"
        "第二级（整体模型）：若分段模型不存在或样本不足，回退至该品牌车系的整体模型（如“本田-飞度”）；\n"
        "第三级（类别模型）：若整体模型仍缺失，回退至该车辆所属类别的通用模型（如“轿车-小型车”）。\n"
        "每个模型均通过竞争优选，在指数函数、二次多项式、幂函数中选择拟合度（R²）最高的作为最终模型。"
    ]
    
    for claim in claims:
        p = doc.add_paragraph(claim)
        set_font(p.runs[0], '宋体', Pt(12))

    # --- Description of Drawings ---
    # 四、附图说明
    head_para = doc.add_paragraph()
    run = head_para.add_run('四、附图说明')
    set_font(run, '黑体', Pt(14), True)
    
    p = doc.add_paragraph(
        "图1 为本发明的方法流程图；\n"
        "图2 为五级特征匹配推断新车价格的逻辑示意图；\n"
        "图3 为自适应分段非线性模型的竞争优选与回退机制示意图。"
    )
    p.paragraph_format.first_line_indent = Inches(0.3)
    set_font(p.runs[0], '宋体', Pt(12))

    # --- Detailed Description ---
    # 五、具体实施方式
    head_para = doc.add_paragraph()
    run = head_para.add_run('五、具体实施方式')
    set_font(run, '黑体', Pt(14), True)
    
    sections = [
        ("1. 多源特征索引构建", 
         "系统首先对原始成交数据进行清洗，去除价格异常点。利用正则表达式和NLP技术构建EntityExtractor，从非结构化的车辆全称（如“大众POLO 2016款 1.6L 自动舒适版”）中提取出：\n"
         "- Year: 2016\n"
         "- Displacement: 1.6L\n"
         "- Transmission: 自动\n"
         "- Trim: 舒适版\n"
         "随后构建ResidualDataIndex，支持按品牌车系、车辆类别及上述细粒度特征的快速检索。"),
         
        ("2. 自适应分段非线性回归建模", 
         "针对二手车残值率随车龄非线性衰减的特性，且不同车龄段衰减速率不同（前期快、后期慢），本系统设计了分段建模策略：\n"
         "将车龄划分为三个区间：Young（0.5-5年）、Mid（5-10年）、Old（10-20年）。\n"
         "对于每个区间，系统分别尝试拟合指数模型、二次多项式模型和幂函数模型。通过计算R²（决定系数）和RMSE（均方根误差），自动选择表现最优的模型并序列化保存。\n"
         "此外，为了解决长尾车型样本不足的问题，系统建立了“分段模型 -> 整体模型 -> 类别模型”的三级兜底机制，确保预测覆盖率达到100%。"),
         
        ("3. 动态置信度修正算法", 
         "模型预测仅反映了历史统计规律，为了捕捉当前市场波动和特定配置的溢价，本发明引入了动态置信度修正（Dynamic Confidence Adjustment）。\n"
         "系统实时检索与待预测车辆特征最相似的Top-K个历史成交案例。对于每个案例，根据其与目标车辆在年款、排量、变速箱、版式、里程、城市等维度的匹配程度，计算相似度得分（Score）。\n"
         "计算所有案例的平均得分作为置信度（Confidence）。若置信度高，说明市场上有大量高度相似的成交案例，预测结果应向这些案例的均价靠拢；若置信度低，则更多依赖模型的统计预测值。")
    ]
    
    for title, content in sections:
        p_title = doc.add_paragraph(title)
        set_font(p_title.add_run(title), '黑体', Pt(12), True)
        
        p_content = doc.add_paragraph(content)
        p_content.paragraph_format.first_line_indent = Inches(0.3)
        set_font(p_content.runs[0], '宋体', Pt(12))

    # Save
    output_path = os.path.join(os.getcwd(), 'output', 'patent_v3.docx')
    if not os.path.exists(os.path.dirname(output_path)):
        os.makedirs(os.path.dirname(output_path))
        
    doc.save(output_path)
    print(f"Patent document v3 generated at: {output_path}")

if __name__ == "__main__":
    create_patent_doc_v3()