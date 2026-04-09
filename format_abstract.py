# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

def set_font(run, font_name, east_asia_font_name):
    """设置字体，区分西文和中文字体"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east_asia_font_name)

doc = docx.Document()

# Set normal style font to ensure fallback works gracefully
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# --- 1. Chinese Title ---
p_title = doc.add_paragraph()
p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run_title = p_title.add_run('基于规则引擎与LightGBM的二手车残值估价系统设计')
set_font(run_title, 'SimHei', '黑体')
run_title.font.size = Pt(16) # 三号
run_title.bold = True
p_title.paragraph_format.space_after = Pt(12)

# --- 2. "摘要" Heading ---
p_abstract_heading = doc.add_paragraph()
p_abstract_heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run_abstract_heading = p_abstract_heading.add_run('摘要')
set_font(run_abstract_heading, 'SimHei', '黑体')
run_abstract_heading.font.size = Pt(16) # 三号
run_abstract_heading.bold = True
p_abstract_heading.paragraph_format.space_after = Pt(12)

# --- 3. Chinese Abstract Body ---
def add_zh_paragraph(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24) # 2 characters for 12pt font
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_font(run, 'Times New Roman', '宋体')
    run.font.size = Pt(12) # 小四

add_zh_paragraph('随着我国二手车市场的快速发展，人们对二手车的精准定价和售卖需求日益突出。然而，由于二手车具有“一车一况、一车一价”的非标准化特性，导致市场信息不对称、定价缺乏透明度与科学性。')
add_zh_paragraph('针对上述问题，本文设计并实现了一套基于规则引擎与LightGBM的二手车残值估价系统。该系统所使用的真实交易数据来源于车易拍、懂车帝等二手车交易平台。在数据预处理与车型对齐阶段，本文设计了一种基于专家规则引擎的车型实体提取与匹配算法。通过构建包含品牌、车系、年款、排量等多维度配置的知识库，结合自定义实体抽取规则与多级相似度计算模型，实现了海量非标准车辆描述到标准车型库的精准映射。在价格预测模块，本文引入了LightGBM机器学习算法，深度挖掘车龄、行驶里程、品牌保值率、城市级别及车辆配置等多维特征。充分利用LightGBM在处理大规模高维数据时的高效性以及捕捉复杂非线性衰减规律的优势，完成了残值率预测模型的训练。系统上线后，能够根据用户输入的车辆多维条件综合评估残值率，从而实现对二手车价格的精准预测。最后，本文采用模块化架构开发了完整的估价应用以及Web测试界面。')

# --- 4. Chinese Keywords ---
p_kw = doc.add_paragraph()
p_kw.paragraph_format.line_spacing = 1.5
run_kw_label = p_kw.add_run('关键词：')
set_font(run_kw_label, 'SimHei', '黑体')
run_kw_label.font.size = Pt(12)
run_kw_label.bold = True

run_kw_text = p_kw.add_run('二手车估价；残值模型；规则引擎；车型匹配；LightGBM；特征工程')
set_font(run_kw_text, 'Times New Roman', '宋体')
run_kw_text.font.size = Pt(12)
p_kw.paragraph_format.space_after = Pt(24)

# --- 5. English Title (Optional but good for thesis) ---
p_en_title = doc.add_paragraph()
p_en_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run_en_title = p_en_title.add_run('Design of Used Car Residual Value Estimation System Based on Rule Engine and LightGBM')
run_en_title.font.name = 'Times New Roman'
run_en_title.font.size = Pt(16)
run_en_title.bold = True
p_en_title.paragraph_format.space_after = Pt(12)

# --- 6. "Abstract" Heading ---
p_en_heading = doc.add_paragraph()
p_en_heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run_en_heading = p_en_heading.add_run('Abstract')
run_en_heading.font.name = 'Times New Roman'
run_en_heading.font.size = Pt(16)
run_en_heading.bold = True
p_en_heading.paragraph_format.space_after = Pt(12)

# --- 7. English Abstract Body ---
def add_en_paragraph(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)

add_en_paragraph('With the rapid development of the used car market in China, the demand for accurate pricing and selling of used cars has become increasingly prominent. However, due to the non-standardized characteristics of used cars—often described as "one condition per car, one price per car"—the market suffers from information asymmetry and a lack of transparency and scientific rigor in pricing.')
add_en_paragraph('To address these issues, this paper designs and implements a used car residual value estimation system based on a rule engine and LightGBM. The real transaction data used in this system is sourced from used car trading platforms such as Cheyipai and Dongchedi. In the data preprocessing and vehicle alignment phase, an expert rule engine-based entity extraction and matching algorithm is designed. By constructing a knowledge base containing multi-dimensional configurations such as brand, series, model year, and engine displacement, combined with custom entity extraction rules and a multi-level similarity calculation model, the precise mapping of massive non-standard vehicle descriptions to a standard vehicle database is achieved. In the price prediction module, the LightGBM machine learning algorithm is introduced to deeply mine multi-dimensional features, including vehicle age, mileage, brand retention rate, city tier, and vehicle configurations. By fully leveraging LightGBM\'s computational efficiency in handling large-scale, high-dimensional data and its advantage in capturing complex non-linear depreciation patterns, the residual value prediction model is successfully trained. Once deployed, the system can comprehensively evaluate the residual value rate based on multiple vehicle conditions input by users, thereby achieving accurate predictions of used car prices. Finally, a complete estimation application and a Web test interface are developed using a modular architecture.')

# --- 8. English Keywords ---
p_en_kw = doc.add_paragraph()
p_en_kw.paragraph_format.line_spacing = 1.5
run_en_kw_label = p_en_kw.add_run('Key Words: ')
run_en_kw_label.font.name = 'Times New Roman'
run_en_kw_label.font.size = Pt(12)
run_en_kw_label.bold = True

run_en_kw_text = p_en_kw.add_run('Used Car Pricing; Residual Value Model; Rule Engine; Vehicle Matching; LightGBM; Feature Engineering')
run_en_kw_text.font.name = 'Times New Roman'
run_en_kw_text.font.size = Pt(12)

doc.save(r'D:\BaiduNetdiskDownload\车型库映射\车型库映射\doc\毕设\基于规则引擎与LightGBM的二手车残值估价系统设计_摘要.docx')
print("Successfully generated beautifully formatted Word document.")
