# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.oxml.ns import qn

doc = docx.Document()

# 基础字体设置
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)  # 小四
style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_run_font(run, ascii_font='Times New Roman', east_asia_font='宋体'):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east_asia_font)

# --- 1. 一级标题 (第一章 绪论) ---
p_title = doc.add_paragraph()
p_title.style = doc.styles['Heading 1'] 
p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
p_title.clear() # 清除默认样式带来的文字
run_title = p_title.add_run('第一章 绪论')
set_run_font(run_title, 'SimHei', '黑体')
run_title.font.size = Pt(16) # 三号
run_title.font.color.rgb = RGBColor(0, 0, 0) # 强制黑色
run_title.bold = True
p_title.paragraph_format.space_before = Pt(12)
p_title.paragraph_format.space_after = Pt(12)
p_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

# --- 2. 二级标题添加函数 ---
def add_heading_2(text):
    p = doc.add_paragraph()
    p.style = doc.styles['Heading 2']
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.clear()
    run = p.add_run(text)
    set_run_font(run, 'SimHei', '黑体')
    run.font.size = Pt(14) # 四号
    run.font.color.rgb = RGBColor(0, 0, 0) # 强制黑色
    run.bold = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

# --- 3. 正文段落添加函数 ---
def add_body_paragraph(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24) # 首行缩进2字符 (对于12pt字体，2个字符=24pt)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE # 绝对1.5倍行距
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, 'Times New Roman', '宋体')
    run.font.size = Pt(12) # 小四
    run.font.color.rgb = RGBColor(0, 0, 0)

# ================= 写入内容 =================

add_heading_2('1.1 课题研究的背景与意义')
add_body_paragraph('随着我国汽车工业的成熟与汽车保有量的不断攀升，汽车消费市场已逐渐由增量市场向存量市场过渡。在这一背景下，二手车市场的交易量呈现出爆发式增长。然而，与标准化的新车交易不同，二手车具有典型的“一车一况、一车一价”特征。车辆的品牌、车系、年款、行驶里程、车况、甚至所在城市等因素，都会对其最终售价产生复杂影响。')
add_body_paragraph('长期以来，国内二手车市场的定价高度依赖于线下评估师的人工经验。这种方式不仅效率低下、人力成本高昂，且不可避免地带有主观随意性。同时，由于各个二手车平台（如车易拍、懂车帝等）的车源描述缺乏统一标准，导致市场信息极度不对称，买卖双方难以建立信任，这成为制约二手车行业规模化、规范化发展的核心痛点。因此，借助大数据与人工智能技术，打破数据孤岛，建立一套客观、透明、高效的二手车残值估价系统，不仅能够降低评估成本、提升平台交易效率，更对推动整个二手车市场的标准化与数字化转型具有重要的现实意义。')

add_heading_2('1.2 课题研究的目的与范围')
add_body_paragraph('本课题旨在设计并开发一套基于规则引擎与LightGBM的二手车残值估价系统。系统的核心目的在于：一方面，通过构建领域知识规则引擎，解决海量异构二手车文本数据的标准化清洗与对齐难题；另一方面，利用前沿的机器学习算法（LightGBM）捕捉二手车残值衰减的复杂非线性规律，提供高精度的自动化估价服务。')
add_body_paragraph('本课题的研究范围涵盖了从数据获取与处理到模型服务上线的完整系统闭环，具体包括：')
add_body_paragraph('（1）数据清洗与特征工程体系设计：收集真实二手车交易数据，设计车龄、行驶里程、品牌保值率、城市级别等多维特征。')
add_body_paragraph('（2）基于规则引擎的车型匹配模块：通过建立包含品牌、车系、年款、排量等多维度的字典库和正则表达式配置，实现非标准化车名到标准车型库的精准映射。')
add_body_paragraph('（3）基于LightGBM的残值估价模型：构建并训练能够适应C2B、B2C等多种业务场景的机器学习定价模型，并与传统的曲线拟合方法进行对比优化。')
add_body_paragraph('（4）Web估价系统与测试平台的实现：提供友好的用户交互界面和API接口，方便系统集成与测试。')

add_heading_2('1.3 拟解决的主要问题')
add_body_paragraph('本课题在系统设计与实现过程中，着重解决以下三个关键问题：')
add_body_paragraph('（1）多源异构车辆描述的标准化消歧问题：二手车市场数据中常常包含错别字、别名、不规范缩写等（例如“21款 比亚迪秦plus dm-i 110km”与标准库名称的差异）。本课题需要通过设计一套多级级联的实体识别与匹配规则引擎，解决非标数据的结构化与精确对齐问题。')
add_body_paragraph('（2）长尾车型与残值非线性衰减的拟合问题：二手车的价值衰减通常呈现“早期陡峭、后期平缓”的非线性规律。传统的简单回归模型或固定曲线公式难以准确刻画，且在样本稀疏的冷门车系（长尾车型）上极易表现不佳或过拟合。本课题需要利用LightGBM树模型的优势来克服这一问题。')
add_body_paragraph('（3）多业务场景的价格矩阵计算问题：不同的交易模式下（如车商收车C2B、车商卖车B2C）车辆的溢价空间不同。系统需要根据不同商业模式设定调整参数，输出符合实际市场规律的多场景价格矩阵。')

add_heading_2('1.4 要求达到的技术参数')
add_body_paragraph('为了确保系统能够在实际业务中稳定且精准地运行，本系统在设计阶段设定了如下技术参数指标与要求：')
add_body_paragraph('（1）车型映射准确率：基于规则引擎的匹配算法，针对带有明显特征描述的二手车文本数据，综合匹配准确率需达到 85% 以上，并在部分高频车系上实现 95% 的匹配成功率。')
add_body_paragraph('（2）残值预测误差率（MAE与RMSE）：基于LightGBM的残值估价模型，其平均绝对误差（MAE）和均方根误差（RMSE）指标需明显优于传统的指数或多项式曲线拟合模型，特别是在长尾数据上需表现出更好的鲁棒性。')
add_body_paragraph('（3）接口响应性能：模型推理的计算延迟应满足生产环境的实时性要求，单次车辆名称映射及价格预测的 API 接口响应时间应控制在 500 毫秒（ms）以内。')
add_body_paragraph('（4）模块化与扩展性：配置文件（YAML）需与核心代码解耦，确保后续新增车型规则或调整模型参数时，无需大规模修改源码即可动态生效。')

doc.save(r'D:\BaiduNetdiskDownload\车型库映射\车型库映射\doc\毕设\第一章 绪论.docx')
print("Successfully generated strictly formatted Chapter 1.")
