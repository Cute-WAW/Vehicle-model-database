# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Title
title = doc.add_heading('基于规则引擎与LightGBM的二手车残值估价系统设计', level=1)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Abstract Heading
abstract_heading = doc.add_heading('摘要', level=2)
abstract_heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Paragraph 1: Background
p1 = doc.add_paragraph()
p1.add_run('随着我国二手车市场的快速发展，二手车交易量逐年攀升。然而，二手车因其“一车一况、一车一价”的非标准化特性，导致市场信息不对称、定价缺乏透明度和科学性。传统的评估方法高度依赖人工经验，不仅效率低下，且难以应对海量且错综复杂的非标准车辆描述数据，以及车辆残值随车龄、品牌、车况等因素呈现的复杂非线性衰减规律。针对上述问题，本研究设计并实现了一套基于规则引擎与LightGBM的二手车残值估价系统，旨在提升二手车车型匹配的准确率与价格评估的科学性。')
p1.paragraph_format.first_line_indent = Pt(21)

# Paragraph 2: Methods
p2 = doc.add_paragraph()
p2.add_run('本研究的主要内容与方法如下：首先，针对多源异构且包含大量非标准描述的二手车交易数据，设计了一种基于专家规则引擎的车型实体提取与匹配算法。通过构建包含品牌、车系、年款、排量、动力类型等多维度配置的知识库，结合自定义实体抽取规则和多级相似度计算模型，实现了海量非标准车名到标准车型库的精准映射。其次，在价格预测模块中，引入LightGBM（轻量级梯度提升机）机器学习算法，替代传统的简单线性回归或曲线拟合模型。在特征工程层面，深度挖掘了车龄、行驶里程、品牌保值率、城市级别及车辆配置等多元特征；在模型构建层面，充分利用LightGBM在处理大规模高维数据时的计算效率和捕捉复杂非线性特征方面的优势，构建了多场景（如C2B、B2C等）二手车残值预测模型。最后，采用模块化架构开发了完整的估价系统以及Web测试界面。')
p2.paragraph_format.first_line_indent = Pt(21)

# Paragraph 3: Results & Conclusions
p3 = doc.add_paragraph()
p3.add_run('研究成果与结论表明：在车型库映射方面，本系统相较于基线方法在匹配准确率和召回率上均有显著提升，能够有效解决长尾车型及模糊描述的对齐难题，准确率达到了业内较高水平；在残值估价方面，基于LightGBM的预测模型有效克服了传统模型在长尾车型上的过拟合及非线性衰减拟合不足等问题，显著降低了平均绝对误差（MAE）和均方根误差（RMSE）。系统整体运行稳定、响应迅速，不仅为二手车电商平台、评估机构等提供了可靠的自动化定价支持，同时为二手车市场的数字化、标准化发展提供了重要的技术参考和应用示范。')
p3.paragraph_format.first_line_indent = Pt(21)

# Keywords
p_kw = doc.add_paragraph()
kw_bold = p_kw.add_run('关键词：')
kw_bold.bold = True
p_kw.add_run('二手车估价；残值模型；规则引擎；车型匹配；LightGBM；特征工程')

doc.save(r'D:\BaiduNetdiskDownload\车型库映射\车型库映射\doc\毕设\基于规则引擎与LightGBM的二手车残值估价系统设计_摘要.docx')
print("Successfully generated Word document.")
